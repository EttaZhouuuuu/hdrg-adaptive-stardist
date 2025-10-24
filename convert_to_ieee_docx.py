"""
Convert markdown to IEEE-style DOCX format
"""

import markdown
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import re

def setup_ieee_styles(doc):
    """Set up IEEE-style formatting"""
    # Title style
    style = doc.styles.add_style('IEEE Title', WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = 'Times New Roman'
    style.font.size = Pt(24)
    style.font.bold = True
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style.paragraph_format.space_after = Pt(12)

    # Heading 1 style
    style = doc.styles.add_style('IEEE Heading 1', WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(8)

    # Normal text style
    style = doc.styles.add_style('IEEE Normal', WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = 'Times New Roman'
    style.font.size = Pt(10)
    style.paragraph_format.space_after = Pt(6)

    # Abstract style
    style = doc.styles.add_style('IEEE Abstract', WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = 'Times New Roman'
    style.font.size = Pt(9)
    style.font.italic = True
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(12)

    # Reference style
    style = doc.styles.add_style('IEEE Reference', WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = 'Times New Roman'
    style.font.size = Pt(8)
    style.paragraph_format.first_line_indent = Inches(-0.25)
    style.paragraph_format.left_indent = Inches(0.25)
    style.paragraph_format.space_after = Pt(3)

def convert_md_to_ieee_docx(input_md, output_docx):
    """Convert markdown file to IEEE-style DOCX"""
    # Read markdown content
    with open(input_md, 'r', encoding='utf-8') as f:
        md_content = f.read()

    # Create new document
    doc = Document()
    
    # Setup IEEE styles
    setup_ieee_styles(doc)

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Split content into sections
    sections = md_content.split('\n## ')
    
    # Process title section
    title_section = sections[0].split('\n\n')
    title = title_section[0].replace('# ', '')
    doc.add_paragraph(title, 'IEEE Title')

    # Add author information
    for line in title_section[1:]:
        if '*IEEE' not in line:  # Skip the format line
            p = doc.add_paragraph(line.strip('*'), 'IEEE Normal')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Process remaining sections
    for section in sections[1:]:
        lines = section.split('\n')
        heading = lines[0]
        content = '\n'.join(lines[1:])

        # Add section heading
        doc.add_paragraph(heading, 'IEEE Heading 1')

        # Process content based on section type
        if 'Abstract' in heading:
            doc.add_paragraph(content.strip(), 'IEEE Abstract')
        elif 'References' in heading:
            # Process each reference
            refs = content.strip().split('\n\n')
            for ref in refs:
                if ref.strip():
                    doc.add_paragraph(ref.strip(), 'IEEE Reference')
        else:
            # Regular content
            paragraphs = content.strip().split('\n\n')
            for para in paragraphs:
                if para.strip():
                    doc.add_paragraph(para.strip(), 'IEEE Normal')

    # Save the document
    doc.save(output_docx)

if __name__ == '__main__':
    input_md = 'AMS_StarDist_IEEE_Paper.md'
    output_docx = 'AMS_StarDist_IEEE_Paper.docx'
    convert_md_to_ieee_docx(input_md, output_docx)
    print(f"Successfully converted {input_md} to {output_docx}")

